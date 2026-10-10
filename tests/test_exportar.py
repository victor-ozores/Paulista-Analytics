import importlib.util
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import verificar_excel  # noqa: E402

PROC = RAIZ / "data" / "processed"
precisa_dados = pytest.mark.skipif(not (PROC / "ref_indicadores.csv").exists(), reason="rode scripts/04_tratar.py e 05_validar.py antes")
precisa_excel = pytest.mark.skipif(not verificar_excel.excel_disponivel(), reason="Excel não encontrado (só Windows com Excel instalado)")

ABAS = ["LEIA-ME", "Resumo", "Classificação", "Artilharia", "Financeiro", "Elenco", "Partidas", "Eventos", "Escalações", "Trocas", "Ingressos", "Despesas", "Rateio", "Clubes", "FPF oficial", "Qualidade", "Log", "Correções", "Dicionário", "Referência Python"]


def carregar_exportar():
    spec = importlib.util.spec_from_file_location("exportar06", RAIZ / "scripts" / "06_exportar.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


@pytest.fixture(scope="module")
def planilha(tmp_path_factory):
    exportar = carregar_exportar()
    destino = tmp_path_factory.mktemp("xlsx") / "conferencia.xlsx"
    return exportar.construir(destino)


@precisa_dados
def test_abas_e_quantidade_de_conferencias(planilha):
    assert planilha["abas"] == ABAS
    wb = load_workbook(planilha["arquivo"])
    assert wb.sheetnames == ABAS
    assert planilha["conferencias"] == 610  # 198 antigas (71 do Resumo + 16 clubes + 111 artilheiros) + 171 indicadores e 72 jogos do Financeiro + 25 indicadores e 144 times-jogo do Elenco
    assert wb["Partidas"].max_row == 73 and wb["Eventos"].max_row == 601 and len(wb["Clubes"]._images) == 16
    assert (wb["Escalações"].max_row, wb["Trocas"].max_row, wb["Ingressos"].max_row, wb["Despesas"].max_row, wb["Rateio"].max_row) == (3293, 707, 2389, 1483, 145)


@precisa_dados
def test_toda_formula_usa_apenas_nomes_definidos(planilha):
    wb = load_workbook(planilha["arquivo"])
    nomes = set(wb.defined_names)
    usados, formulas = set(), 0
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    sem_textos = re.sub(r'"[^"]*"', "", c.value)
                    usados |= set(re.findall(r"\b(?:P|E|O|R|Q|L|X|IG|DS|RT|ES|TR|Par|Conf)_[A-Za-z0-9_]+\b", sem_textos))
    assert formulas > 1000
    assert usados <= nomes, f"nomes usados em fórmulas e não definidos: {sorted(usados - nomes)}"
    assert {"Conf_total", "Conf_ok", "Conf_status_resumo", "Conf_status_class", "Conf_status_art", "Conf_status_fin", "Conf_status_fin_jogo", "Conf_status_elenco", "Conf_status_elenco_tj", "Par_fase"} <= nomes


@precisa_dados
def test_nomes_definidos_apontam_para_dentro_da_aba(planilha):
    wb = load_workbook(planilha["arquivo"])
    for nome, dn in wb.defined_names.items():
        for aba, ref in dn.destinations:
            ws = wb[aba]
            celulas = ws[ref.replace("$", "")]
            ultima = celulas[-1][-1] if isinstance(celulas, tuple) and isinstance(celulas[-1], tuple) else celulas
            assert ultima.row <= ws.max_row, f"{nome} passa do fim da aba {aba}"


@precisa_dados
def test_dicionario_tem_os_59_indicadores(planilha):
    ws = load_workbook(planilha["arquivo"])["Dicionário"]
    ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    assert ids == [f"I{i:02d}" for i in range(1, 60)]


@precisa_dados
@precisa_excel
def test_no_excel_de_verdade_tudo_confere_e_defeito_plantado_e_detectado(planilha, tmp_path):
    res = verificar_excel.verificar(planilha["arquivo"])
    assert res["erros"] == [] and res["fora_do_ok"] == []
    assert res["total"] == res["ok"] == planilha["conferencias"]
    wb = load_workbook(planilha["arquivo"])
    ws = wb["Partidas"]
    cab = {c.value: c.column for c in ws[1]}
    ws.cell(2, cab["Gols mandante"]).value += 1  # muda um placar de propósito
    adulterada = tmp_path / "adulterada.xlsx"
    wb.save(adulterada)
    res = verificar_excel.verificar(adulterada)
    assert res["erros"] == [] and len(res["fora_do_ok"]) >= 5
    assert res["ok"] < res["total"]


@precisa_dados
def test_financeiro_confere_os_indicadores_i38_a_i54_por_escopo_e_cada_um_dos_72_jogos(planilha):
    ws = load_workbook(planilha["arquivo"])["Financeiro"]
    indicadores = [(ws.cell(r, 1).value, ws.cell(r, 3).value) for r in range(6, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, str) and re.fullmatch(r"I\d\d", ws.cell(r, 1).value)]
    assert len(indicadores) == 171 and {i for i, _ in indicadores} == {f"I{n}" for n in range(38, 55)}
    escopos = {}
    for _, esc in indicadores:
        escopos[esc] = escopos.get(esc, 0) + 1
    assert escopos == {"geral": 7, "fase": 28, "categoria": 30, "grupo_despesa": 20, "bloco": 6, "clube": 80}
    assert all(str(ws.cell(r, 5).value).startswith("=") and str(ws.cell(r, 6).value).startswith("=INDEX(R_valor") for r in range(6, 6 + 171))
    por_jogo = [r for r in range(6 + 171, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)]
    assert [ws.cell(r, 1).value for r in por_jogo] == list(range(1, 73)) and all(str(ws.cell(r, 12).value).startswith("=IF(AND(") for r in por_jogo)


@precisa_dados
def test_elenco_confere_os_indicadores_i55_a_i59_e_cada_time_em_cada_jogo(planilha):
    ws = load_workbook(planilha["arquivo"])["Elenco"]
    assert [ws.cell(r, 1).value for r in range(6, 31)] == ["I55", "I56", "I57", "I58", "I59"] * 5
    assert [ws.cell(r, 5).value for r in range(6, 11)] == ["=COUNTIF(ES_jogou,1)", '=COUNTIF(ES_situacao,"titular")', '=COUNTIFS(ES_situacao,"reserva",ES_jogou,1)', "=ROWS(TR_jogo)", '=SUMIFS(ES_gols,ES_situacao,"reserva")']
    times = [r for r in range(31, ws.max_row + 1) if ws.cell(r, 3).value in ("M", "V")]
    assert len(times) == 144 and [ws.cell(r, 3).value for r in times[:4]] == ["M", "V", "M", "V"] and {ws.cell(r, 1).value for r in times} == set(range(1, 73))
    assert str(ws.cell(times[0], 15).value).startswith("=IF(AND(E") and "H" in ws.cell(times[0], 15).value


@precisa_dados
def test_novas_abas_de_dados_trazem_as_linhas_das_tabelas_tratadas(planilha):
    wb = load_workbook(planilha["arquivo"])
    soma = lambda aba, coluna: sum(wb[aba].cell(r, coluna).value or 0 for r in range(2, wb[aba].max_row + 1))
    assert soma("Ingressos", 8) == 784369 and round(soma("Ingressos", 10), 2) == 44730003.80
    assert round(soma("Despesas", 7), 2) == 20404747.14
    assert round(soma("Rateio", 9), 2) == 24171721.26 and round(soma("Rateio", 7) + soma("Rateio", 8), 2) == 153535.40
    assert sum(1 for r in range(2, 3294) if wb["Escalações"].cell(r, 6).value == "titular") == 1584 and soma("Escalações", 8) == 2290
    assert soma("Escalações", 12) == 166 and soma("Escalações", 13) == 2, "166 gols de atletas e 2 gols contra"
    assert sum(1 for r in range(2, 708) if wb["Trocas"].cell(r, 5).value == "concussao") == 2


@precisa_dados
@precisa_excel
@pytest.mark.parametrize("aba,coluna,novo", [
    ("Ingressos", "Ingressos", lambda v: v + 1),
    ("Despesas", "Valor", lambda v: v + 1),
    ("Rateio", "Líquido a receber", lambda v: v + 1),
    ("Escalações", "Jogou (1/0)", lambda v: 0),
    ("Escalações", "Gols", lambda v: v + 1),
    ("Trocas", "Clube", lambda v: "Palmeiras" if v != "Palmeiras" else "Santos"),
])
def test_no_excel_defeito_plantado_nas_novas_abas_e_detectado(planilha, tmp_path, aba, coluna, novo):
    wb = load_workbook(planilha["arquivo"])
    ws = wb[aba]
    cab = {c.value: c.column for c in ws[1]}
    celula = ws.cell(2, cab[coluna])
    celula.value = novo(celula.value)
    adulterada = tmp_path / "adulterada.xlsx"
    wb.save(adulterada)
    res = verificar_excel.verificar(adulterada)
    assert res["erros"] == [] and len(res["fora_do_ok"]) >= 1, (aba, coluna)
    assert res["ok"] < res["total"]

